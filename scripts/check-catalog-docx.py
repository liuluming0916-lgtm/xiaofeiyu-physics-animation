"""校验重写后的考点 docx：表格数、考点数、有没有空单元格。"""
from docx import Document

P = "开源日志/人教版高中物理-动画考点提示词-重写版.docx"
doc = Document(P)

print("表格数:", len(doc.tables), "段落数:", len(doc.paragraphs))
rows = 0
empty = 0
first = None
for t in doc.tables:
    for r in t.rows[1:]:
        cells = [c.text.strip() for c in r.cells]
        rows += 1
        if any(not c for c in cells):
            empty += 1
        if first is None:
            first = cells

print("考点行数:", rows, "空单元格行:", empty)
print("首行考点:", first[0] if first else "-")
print("提示词是否含英文 Tags:", any("English Tags" in c.text for t in doc.tables for r in t.rows for c in r.cells))
print("提示词是否含分屏/双栏字样:", any(
    ("分屏" in c.text or "双栏" in c.text or "2.5D" in c.text or "Manim" in c.text)
    for t in doc.tables for r in t.rows[1:] for c in r.cells
))
