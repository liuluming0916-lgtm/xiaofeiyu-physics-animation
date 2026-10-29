"""把 Gemini 生成的考点 docx 抽成 markdown，便于查看结构（输出到 开源日志/_check）。"""
import os
from docx import Document

SRC = r"C:\Users\Liu LuMing\Desktop\各种md\docs\人教版高中物理教材可动画化考点提示词大全.docx"
OUT = r"E:\XFY-2.0\开源日志\_check\catalog-dump.md"

doc = Document(SRC)
lines = []

body = doc.element.body
para_iter = iter(doc.paragraphs)
table_iter = iter(doc.tables)

# 按文档顺序遍历段落与表格
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

for child in body.iterchildren():
    if child.tag == qn("w:p"):
        p = Paragraph(child, doc)
        t = p.text.strip()
        if t:
            lines.append(f"## {t}" if len(t) < 60 else t)
    elif child.tag == qn("w:tbl"):
        tb = Table(child, doc)
        lines.append("")
        for row in tb.rows:
            cells = [c.text.strip().replace("\n", " / ") for c in row.cells]
            lines.append("| " + " | ".join(cells) + " |")
        lines.append("")

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    f.write("\n".join(lines))

print("段落", len(doc.paragraphs), "表格", len(doc.tables), "输出行", len(lines))
