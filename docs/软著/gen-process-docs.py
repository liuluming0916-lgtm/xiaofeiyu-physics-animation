"""把 docs/软著/过程材料/*.md 转成同名的 .docx（软著过程材料用）。

支持：# 标题 / - 列表 / 1. 编号 / | 表格 | / ``` 代码块 / 普通段落。
"""

import re
import sys
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

SRC = Path(__file__).resolve().parent / "过程材料"

BODY_CN = "宋体"
BODY_EN = "Times New Roman"
HEAD_CN = "黑体"
CODE_EN = "Courier New"
BODY_SIZE = Pt(10.5)


def set_run(run, *, cn=BODY_CN, en=BODY_EN, size=BODY_SIZE, bold=False):
    run.font.name = en
    run.font.size = size
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), cn)
    rfonts.set(qn("w:ascii"), en)
    rfonts.set(qn("w:hAnsi"), en)


def new_doc() -> Document:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21)
    sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.2)
    sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.4)
    sec.right_margin = Cm(2.2)
    style = doc.styles["Normal"]
    style.font.name = BODY_EN
    style.font.size = BODY_SIZE
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), BODY_CN)
    style.paragraph_format.space_after = Pt(4)
    style.paragraph_format.line_spacing = 1.35
    return doc


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(text)
    size = Pt(16) if level == 1 else Pt(13) if level == 2 else Pt(11.5)
    set_run(run, cn=HEAD_CN, en=BODY_EN, size=size, bold=True)


INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_runs(p, text, *, size=BODY_SIZE, bold=False):
    """把行内 markdown 变成真正的格式：**粗体** → 粗体，`代码` → 等宽。

    不解析的话，`**` 会原样印进 Word（v1.2.0 文档里曾出现 70+ 处），
    看起来像没写完的草稿。
    """
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            run = p.add_run(part[2:-2])
            set_run(run, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`") and len(part) > 2:
            run = p.add_run(part[1:-1])
            set_run(run, size=size, bold=bold, en=CODE_EN)
        else:
            run = p.add_run(part)
            set_run(run, size=size, bold=bold)


def add_bullet(doc, text, numbered=False):
    p = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
    p.paragraph_format.space_after = Pt(2)
    add_runs(p, text)


def add_para(doc, text):
    p = doc.add_paragraph()
    add_runs(p, text)


def add_code(doc, lines):
    for line in lines:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line if line else " ")
        set_run(run, cn=BODY_CN, en=CODE_EN, size=Pt(9))


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    for i, cell in enumerate(table.rows[0].cells):
        cell.text = ""
        add_runs(cell.paragraphs[0], header[i], size=Pt(9.5), bold=True)
    for row in body:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            if i >= len(cells):
                break
            cells[i].text = ""
            add_runs(cells[i].paragraphs[0], text, size=Pt(9.5))


def split_row(line):
    inner = line.strip().strip("|")
    return [c.strip() for c in inner.split("|")]


def is_sep(line):
    return bool(re.fullmatch(r"\|[\s:\-|]+\|", line.strip()))


def convert(md_path: Path) -> Path:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = new_doc()

    i = 0
    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        if stripped.startswith("```"):
            fence, i = [], i + 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                fence.append(lines[i])
                i += 1
            add_code(doc, fence)
            i += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                if not is_sep(lines[i]):
                    rows.append(split_row(lines[i]))
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        if not stripped:
            i += 1
            continue

        if stripped.startswith("### "):
            add_heading(doc, stripped[4:], 3)
        elif stripped.startswith("## "):
            add_heading(doc, stripped[3:], 2)
        elif stripped.startswith("# "):
            add_heading(doc, stripped[2:], 1)
        elif re.match(r"^\d+\.\s", stripped):
            add_bullet(doc, re.sub(r"^\d+\.\s", "", stripped), numbered=True)
        elif stripped.startswith("- "):
            add_bullet(doc, stripped[2:])
        elif stripped.startswith("> "):
            add_para(doc, stripped[2:])
        else:
            add_para(doc, stripped)
        i += 1

    out = md_path.with_suffix(".docx")
    doc.save(out)
    return out


def main() -> None:
    sys.stdout.reconfigure(encoding="utf-8")
    files = sorted(SRC.glob("*.md"))
    if not files:
        print("没找到 md 文件")
        return
    for md in files:
        out = convert(md)
        print(f"✓ {out.name} · {out.stat().st_size / 1024:.0f} KB")
    print(f"共 {len(files)} 份")


if __name__ == "__main__":
    main()
