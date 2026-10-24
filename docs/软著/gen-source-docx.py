"""生成软著提交用的源代码文档（前 30 页 + 后 30 页，每页 50 行）。

规矩（中国版权保护中心口径）：
  · 每页 50 行代码（不含页眉页码）
  · 页眉写「软件名称 + 版本号」，须与申请表逐字一致
  · 页码连续
  · 总页数超过 60 页时，只交前 30 页与后 30 页

输出：docs/软著/源代码-提交版.docx
"""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

APP_NAME = "小肥鱼物理解题动画软件"
APP_VERSION = "V1.0"

LINES_PER_PAGE = 50
PAGES_EACH_SIDE = 30
MAX_UNITS = 100
FONT = "Courier New"
FONT_SIZE = Pt(8)
LINE_HEIGHT = Pt(12)

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "软著" / "源代码-提交版.docx"

# 拼接顺序：前 30 页要能看出主链路，后 30 页收在界面层
SOURCES = [
    "core/orchestrator/pipeline.mjs",
    "core/analysis/generate.mjs",
    "core/lint/contract-lint.mjs",
    "core/lint/lint-slot.mjs",
    "core/dryrun/dry-run.mjs",
    "core/dryrun/worker-sandbox.mjs",
    "core/repair/classify.mjs",
    "core/export/page.mjs",
    "core/export/gif.mjs",
    "electron/main.cjs",
    "web/src/pages/SolvePage.tsx",
    "web/src/pages/AnimationPage.tsx",
    "web/src/pages/SettingsPage.tsx",
    "web/src/App.tsx",
]


def char_units(ch: str) -> int:
    """等宽字体下按「半角=1、全角=2」计宽（中文注释是主要来源）。"""
    return 2 if ord(ch) > 0x2E7F else 1


def expand(line: str) -> list[str]:
    """制表符转空格；按**显示宽度**折行，保证一页 50 行不被自动折行打乱。"""
    line = line.replace("\t", "    ").rstrip()
    if sum(char_units(c) for c in line) <= MAX_UNITS:
        return [line]
    out = []
    buf = ""
    width = 0
    for ch in line:
        w = char_units(ch)
        if width + w > MAX_UNITS:
            out.append(buf)
            buf = "    " + ch
            width = 4 + w
        else:
            buf += ch
            width += w
    out.append(buf)
    return out


def collect() -> list[str]:
    lines: list[str] = []
    used = []
    for rel in SOURCES:
        path = ROOT / rel
        if not path.exists():
            continue
        used.append(rel)
        lines.append(f"// ===== {rel} =====")
        text = path.read_text(encoding="utf-8")
        for raw in text.splitlines():
            lines.extend(expand(raw))
        lines.append("")
    return lines, used


def set_base_style(doc: Document) -> None:
    style = doc.styles["Normal"]
    style.font.name = FONT
    style.font.size = FONT_SIZE
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:eastAsia"), "SimSun")
    pf = style.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = LINE_HEIGHT


def add_page_number(paragraph) -> None:
    paragraph.alignment = 1
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)


def main() -> None:
    lines, used = collect()
    total_pages = (len(lines) + LINES_PER_PAGE - 1) // LINES_PER_PAGE

    if total_pages > PAGES_EACH_SIDE * 2:
        head = lines[: PAGES_EACH_SIDE * LINES_PER_PAGE]
        tail = lines[-PAGES_EACH_SIDE * LINES_PER_PAGE :]
        lines = head + tail
        mode = f"前 {PAGES_EACH_SIDE} 页 + 后 {PAGES_EACH_SIDE} 页"
    else:
        mode = "全部"

    pages = [lines[i : i + LINES_PER_PAGE] for i in range(0, len(lines), LINES_PER_PAGE)]
    for page in pages:
        while len(page) < LINES_PER_PAGE:
            page.append("")

    doc = Document()
    set_base_style(doc)

    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.6)
    section.left_margin = Cm(1.9)
    section.right_margin = Cm(1.6)

    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = 1
    hr = header.add_run(f"{APP_NAME} {APP_VERSION}")
    hr.font.name = FONT
    hr.font.size = Pt(9)

    add_page_number(section.footer.paragraphs[0])

    for index, page in enumerate(pages):
        for offset, line in enumerate(page):
            p = doc.add_paragraph()
            run = p.add_run(line)
            run.font.name = FONT
            run.font.size = FONT_SIZE
            if index < len(pages) - 1 and offset == LINES_PER_PAGE - 1:
                run.add_break(WD_BREAK.PAGE)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)

    print(f"来源文件 {len(used)} 个")
    for rel in used:
        print(f"  · {rel}")
    print(f"原始总页数 {total_pages} → 提交 {len(pages)} 页（{mode}）")
    print(f"写出 {OUT.relative_to(ROOT)} · {OUT.stat().st_size / 1024:.0f} KB")


if __name__ == "__main__":
    main()
